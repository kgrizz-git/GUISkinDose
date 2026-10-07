"""Phase 5 tests for the HTML report writer + format dispatcher."""

from __future__ import annotations

import base64

import pandas as pd

from guiskindose import PyskindoseSettings, load_settings_example_json
from guiskindose.export import ExportExamSource, ExportSource, collect_export_payload
from guiskindose.export.models import ImageEntry
from guiskindose.export.writers import FORMATS, render_bytes
from guiskindose.export.writers.html import render_html_bytes

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _payload():
    out = {
        "psd": 1.0, "air_kerma": 3.0,
        "patient": {"patient_type": "human", "patient": {
            "human_phantom": "hudfrid",
            "patient_skin_cells": {"x": [0.0, 1.0, 2.0], "y": [0.0, 0.0, 0.0], "z": [0.0, 0.0, 0.0]},
            "triangle_vertex_indices": {"i": [0], "j": [1], "k": [2]}}, "orientation": "hfs", "offsets": {}},
        "dose_map": [[1, 1.0]],
        "corrections": {"correction_value_index": [[0, 1]], "backscatter": [[1.2, 1.4]],
                        "inverse_square_law": [[0.9, 0.8]], "medium": [1.02], "table": [0.8], "kerma": [3.0]},
    }
    s = PyskindoseSettings(settings=load_settings_example_json(), output_format="dict")
    src = ExportSource(execution_context="cli", output_dict=out,
                       exams=[ExportExamSource("e1", pd.DataFrame(), None, "e1.dcm", s, (0, 0, 0))],
                       file_name="e1.dcm")
    return collect_export_payload(src, with_images=False)


def test_write_html_smoke():
    payload = _payload()
    payload.images = [ImageEntry(label="Dorsal", view="dorsal", exam_id=None, png_bytes=_PNG)]
    html = render_html_bytes(payload).decode()
    assert "<html" in html
    assert "data:image/png;base64," in html  # embedded image
    assert "schema_version" in html


def test_write_html_missing_image_notice():
    payload = _payload()
    payload.images = [ImageEntry(label="Dorsal", view="dorsal", exam_id=None, png_bytes=None,
                                 error_message="Image unavailable (kaleido/export error)")]
    html = render_html_bytes(payload).decode()
    assert "Image unavailable" in html


def test_dispatcher_all_formats():
    import pytest

    payload = _payload()
    for fmt in FORMATS:
        if fmt == "pdf":
            pytest.importorskip("reportlab")  # optional `export` extra
        data = render_bytes(payload, fmt)
        assert isinstance(data, bytes)
        assert len(data) > 100


def test_html_report_carries_intended_use_notice():
    from guiskindose.intended_use import INTENDED_USE_NOTICE

    html = render_html_bytes(_payload()).decode()
    assert "Intended use:" in html
    assert "not FDA-cleared" in html
    assert _payload().intended_use == INTENDED_USE_NOTICE


_TUBE_SUMMARY = [
    {"tube": "A", "events": 2, "kerma_reported": 3.0, "kerma_corrected": 3.3, "applied_cf": 1.1, "peak_dose": 4.0},
    {"tube": "B", "events": 1, "kerma_reported": 2.0, "kerma_corrected": 2.0, "applied_cf": 1.0, "peak_dose": 0.0},
]


def _tube_payload(summary):
    payload = _payload()
    payload.exams[0].tube_summary = summary
    return payload


def test_all_formats_carry_dose_by_tube_section():
    """HTML, XLSX, DOCX and PDF all show the per-tube table and the non-additive note."""
    import io

    import pytest
    from openpyxl import load_workbook

    from guiskindose.export.sections import TUBE_NOTE

    payload = _tube_payload(_TUBE_SUMMARY)
    html = render_html_bytes(payload).decode()
    assert "Dose by tube" in html
    assert "Plane A" in html
    assert "Plane B" in html
    assert TUBE_NOTE[:30] in html

    wb = load_workbook(io.BytesIO(render_bytes(payload, "xlsx")))
    assert "Dose by tube" in wb.sheetnames

    docx = pytest.importorskip("docx")
    doc = docx.Document(io.BytesIO(render_bytes(payload, "docx")))
    assert any("Dose by tube" in p.text for p in doc.paragraphs)

    pytest.importorskip("reportlab")
    assert render_bytes(payload, "pdf").startswith(b"%PDF")


def test_single_tube_and_missing_summary_add_no_section():
    import io

    from openpyxl import load_workbook

    for summary in (None, [{**_TUBE_SUMMARY[0], "tube": "single"}]):
        payload = _tube_payload(summary)
        assert "Dose by tube" not in render_html_bytes(payload).decode()
        assert "Dose by tube" not in load_workbook(io.BytesIO(render_bytes(payload, "xlsx"))).sheetnames


def test_exam_heading_does_not_double_the_exam_prefix():
    from guiskindose.export._format import exam_heading

    assert exam_heading("Exam 1") == "Exam 1"
    assert exam_heading("e1") == "Exam e1"
