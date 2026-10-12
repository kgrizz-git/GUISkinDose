"""Tests for named XLSX sheet selection in the canonical profile seam.

Sheet names resolve inside the guarded adapter against the opened workbook's
own catalog on the same bounded bytes used for hashing and parsing — there is
no separate listing pass. Covers name/index/default equivalence, unknown and
duplicate titles failing closed without echo, preflight rejection before any
workbook open, and single-snapshot reads. Workbooks are generated in memory
(or under ``tmp_path``) so no tracked ``.xlsx`` binary is ever committed.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import pytest

from guiskindose import get_path_to_example_rdsr_files
from guiskindose.__main__ import cli
from guiskindose.cli_equipment_profiles import equipment_session_from_args
from guiskindose.equipment_profiles import (
    FileSource,
    active_set_hash,
    assemble_session,
)
from guiskindose.equipment_profiles.constants import MAX_SOURCE_BYTES
from guiskindose.equipment_profiles.errors import EquipmentProfileError, ProfileErrorCode
from guiskindose.equipment_profiles.xlsx_source import parse_xlsx_bytes
from tests.unittests.test_cli_equipment_profiles import (
    _COLUMNS,
    _argv_namespace,
    _dose_argv,
    _settings_file,
    _siemens_profile_doc,
    _siemens_table,
    _stdout_json,
    _write_xlsx,
)

_SIEMENS_DICOM = get_path_to_example_rdsr_files() / "siemens_axiom_artis.dcm"

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import Workbook  # noqa: E402

_SIEMENS_LEAVES = [
    ("translation_offset.x", "5.0"),
    ("translation_offset.y", "-3.5"),
    ("translation_offset.z", "0.0"),
    ("translation_direction.x", "-"),
    ("translation_direction.y", "+"),
    ("translation_direction.z", "+"),
    ("rotation_direction.Ap1", "-"),
    ("rotation_direction.Ap2", "+"),
    ("rotation_direction.Ap3", "+"),
    ("rotation_direction.At1", "+"),
    ("rotation_direction.At2", "-"),
    ("rotation_direction.At3", "+"),
    ("swap_lateral_longitudinal", "false"),
    ("field_size_mode", "CFA"),
    ("detector_side_length", "30.0"),
]

_PHILIPS_LEAVES = [
    ("translation_offset.x", "1.0"),
    ("translation_offset.y", "2.0"),
    ("translation_offset.z", "3.0"),
    ("translation_direction.x", "+"),
    ("translation_direction.y", "+"),
    ("translation_direction.z", "+"),
    ("rotation_direction.Ap1", "+"),
    ("rotation_direction.Ap2", "-"),
    ("rotation_direction.Ap3", "+"),
    ("rotation_direction.At1", "+"),
    ("rotation_direction.At2", "+"),
    ("rotation_direction.At3", "-"),
    ("swap_lateral_longitudinal", "true"),
    ("field_size_mode", "ASD"),
    ("detector_side_length", "25.0"),
]


def _table(
    profile_id: str, manufacturer: str, model: str, leaves: list[tuple[str, str]]
) -> list[list[str]]:
    head = [""] * len(_COLUMNS)
    head[0] = "profile"
    head[1] = profile_id
    head[_COLUMNS.index("manufacturer")] = manufacturer
    head[_COLUMNS.index("models")] = model
    rows = [head]
    for field, value in leaves:
        row = [""] * len(_COLUMNS)
        row[0] = "coordinate"
        row[1] = profile_id
        row[_COLUMNS.index("field")] = field
        row[_COLUMNS.index("value")] = value
        rows.append(row)
    return rows


def _philips_table() -> list[list[str]]:
    return _table("p-philips", "Philips", "Allura Clarity", _PHILIPS_LEAVES)


def _workbook_bytes(sheets: list[tuple[str, list[list[str]]]]) -> bytes:
    """Build a workbook with titled sheets from string rows (no tracked binary)."""
    import io

    workbook = Workbook()
    first = True
    for title, rows in sheets:
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        assert sheet is not None
        sheet.title = title
        sheet.append(list(_COLUMNS))
        for row in rows:
            sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _two_sheet_bytes() -> bytes:
    return _workbook_bytes(
        [
            ("SiemensSheet", _table("p-siemens", "Siemens", "AXIOM-Artis", _SIEMENS_LEAVES)),
            ("PhilipsSheet", _table("p-philips", "Philips", "Allura Clarity", _PHILIPS_LEAVES)),
        ]
    )


def _code(excinfo: pytest.ExceptionInfo[EquipmentProfileError]) -> ProfileErrorCode:
    return excinfo.value.code


def test_named_equals_indexed_and_default() -> None:
    raw = _two_sheet_bytes()
    named = parse_xlsx_bytes(raw, sheet="PhilipsSheet")
    indexed = parse_xlsx_bytes(raw, sheet=1)
    assert active_set_hash(named.profiles) == active_set_hash(indexed.profiles)
    assert [profile.profile_id for profile in named.profiles] == ["p-philips"]
    defaulted = parse_xlsx_bytes(raw)
    first = parse_xlsx_bytes(raw, sheet=0)
    assert active_set_hash(defaulted.profiles) == active_set_hash(first.profiles)
    assert [profile.profile_id for profile in defaulted.profiles] == ["p-siemens"]


def test_file_source_name_matches_index(tmp_path: Path) -> None:
    target = tmp_path / "profiles.xlsx"
    target.write_bytes(_two_sheet_bytes())
    named = assemble_session(files=[FileSource(str(target), sheet="PhilipsSheet")])
    indexed = assemble_session(files=[FileSource(str(target), sheet=1)])
    assert named.active_set_hash == indexed.active_set_hash
    assert [source.profile_ids for source in named.provenance.sources] == [("p-philips",)]
    assert "PhilipsSheet" not in repr(named.provenance)


def test_unknown_name_is_value_free() -> None:
    raw = _two_sheet_bytes()
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(raw, sheet="SENTINEL_MISSING_TITLE")
    assert _code(excinfo) is ProfileErrorCode.UNKNOWN_SHEET
    assert "SENTINEL_MISSING_TITLE" not in str(excinfo.value)
    assert "SENTINEL_MISSING_TITLE" not in repr(excinfo.value)


def test_duplicate_titles_fail_closed() -> None:
    import io
    import zipfile

    raw = _two_sheet_bytes()
    forged = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as archive, zipfile.ZipFile(forged, "w") as writer:
        for info in archive.infolist():
            data = archive.read(info.filename)
            if info.filename == "xl/workbook.xml":
                data = data.replace(b"PhilipsSheet", b"SiemensSheet")
            writer.writestr(info, data)
    forged_bytes = forged.getvalue()
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(forged_bytes, sheet="SiemensSheet")
    assert _code(excinfo) is ProfileErrorCode.UNKNOWN_SHEET
    assert "SiemensSheet" not in str(excinfo.value)


def test_duplicate_titles_ambiguous_without_forgery() -> None:
    from guiskindose.equipment_profiles.xlsx_preflight import resolve_sheet

    with pytest.raises(EquipmentProfileError) as excinfo:
        resolve_sheet(2, "Same", titles=["Same", "Same"])
    assert excinfo.value.code is ProfileErrorCode.UNKNOWN_SHEET


def _assert_pre_open_rejection(raw: bytes, monkeypatch: pytest.MonkeyPatch) -> None:
    opens: list[bool] = []
    real_load = openpyxl.load_workbook

    def _spy_open(*args: Any, **kwargs: Any) -> Any:
        opens.append(True)
        return real_load(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", _spy_open)
    with pytest.raises(EquipmentProfileError):
        parse_xlsx_bytes(raw, sheet="Anything")
    assert opens == []


def test_preflight_rejects_before_open(monkeypatch: pytest.MonkeyPatch) -> None:
    _assert_pre_open_rejection(b"not-a-workbook", monkeypatch)
    _assert_pre_open_rejection(b"x" * (MAX_SOURCE_BYTES + 1), monkeypatch)
    _assert_pre_open_rejection(_two_sheet_bytes()[: len(_two_sheet_bytes()) // 2], monkeypatch)


def test_sheet_count_rejects_before_open(monkeypatch: pytest.MonkeyPatch) -> None:
    import io

    workbook = Workbook()
    for index in range(17):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        assert sheet is not None
        sheet.title = f"Sheet{index}"
        sheet.append(["record_type"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    opens: list[bool] = []
    real_load = openpyxl.load_workbook

    def _spy_open(*args: Any, **kwargs: Any) -> Any:
        opens.append(True)
        return real_load(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", _spy_open)
    raised_argument_1 = buffer.getvalue()
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(raised_argument_1, sheet="Sheet16")
    assert _code(excinfo) is ProfileErrorCode.SHEET_LIMIT_EXCEEDED
    assert opens == []


def test_named_read_uses_single_snapshot(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import guiskindose.equipment_profiles.xlsx_source as xlsx_source

    target = tmp_path / "profiles.xlsx"
    target.write_bytes(_two_sheet_bytes())
    reads: list[str] = []
    real_read = xlsx_source._bounded_read
    replacement = _workbook_bytes(
        [("PhilipsSheet", _table("p-philips", "Philips", "Allura Clarity", _PHILIPS_LEAVES))]
    )

    def _spy_read(path: Any, source_kind: str) -> bytes:
        reads.append(str(path))
        raw = real_read(path, source_kind)
        Path(path).write_bytes(replacement)
        return raw

    monkeypatch.setattr(xlsx_source, "_bounded_read", _spy_read)
    session = assemble_session(files=[FileSource(str(target), sheet="SiemensSheet")])
    assert reads == [str(target)]
    assert [source.profile_ids for source in session.provenance.sources] == [("p-siemens",)]


def test_workbook_opened_once_per_parse(monkeypatch: pytest.MonkeyPatch) -> None:
    opens: list[bool] = []
    real_load = openpyxl.load_workbook

    def _spy_open(*args: Any, **kwargs: Any) -> Any:
        opens.append(True)
        return real_load(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", _spy_open)
    parse_xlsx_bytes(_two_sheet_bytes(), sheet="PhilipsSheet")
    assert len(opens) == 1


# ---------------------------------------------------------------------------
# CLI sheet selection through the same seam
# ---------------------------------------------------------------------------


def test_sheet_named_and_indexed(tmp_path: Path) -> None:
    workbook = tmp_path / "profiles.xlsx"
    _write_xlsx(
        workbook,
        [
            ("SiemensSheet", list(_COLUMNS), _siemens_table()),
            ("PhilipsSheet", list(_COLUMNS), _philips_table()),
        ],
    )
    named = equipment_session_from_args(
        _argv_namespace([workbook], sheet="PhilipsSheet"),
    )
    assert named is not None
    assert [source.profile_ids for source in named.provenance.sources] == [("p-philips",)]
    indexed = equipment_session_from_args(_argv_namespace([workbook], sheet="1"))
    assert indexed is not None
    assert indexed.active_set_hash == named.active_set_hash
    first = equipment_session_from_args(_argv_namespace([workbook], sheet="0"))
    assert first is not None
    assert [source.profile_ids for source in first.provenance.sources] == [("p-siemens",)]
    defaulted = equipment_session_from_args(_argv_namespace([workbook]))
    assert defaulted is not None
    assert defaulted.active_set_hash == first.active_set_hash


def test_sheet_dose_parity(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    workbook = tmp_path / "profiles.xlsx"
    _write_xlsx(
        workbook,
        [
            ("SiemensSheet", list(_COLUMNS), _siemens_table()),
            ("PhilipsSheet", list(_COLUMNS), _philips_table()),
        ],
    )
    settings_path = _settings_file(tmp_path)
    monkeypatch.setattr(
        sys,
        "argv",
        _dose_argv(settings_path, _SIEMENS_DICOM, [workbook], ["--equipment-profile-sheet", "SiemensSheet"]),
    )
    cli()
    payload = _stdout_json(capsys)
    assert payload["equipment_profile_match"]["profile_id"] == "p-siemens"


def test_cli_unknown_sheet_and_meaningless_selector(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    workbook = tmp_path / "profiles.xlsx"
    _write_xlsx(workbook, [("Profiles", list(_COLUMNS), _siemens_table())])
    json_profile = tmp_path / "profile.json"
    json_profile.write_text(json.dumps(_siemens_profile_doc()), encoding="utf-8")
    settings_path = _settings_file(tmp_path)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "guiskindose",
            "--settings",
            str(settings_path),
            "--file-path",
            str(_SIEMENS_DICOM),
            "--equipment-profile",
            str(workbook),
            "--equipment-profile-sheet",
            "SENTINEL_SHEET_NAME",
        ],
    )
    with pytest.raises(SystemExit) as excinfo:
        cli()
    assert excinfo.value.code == 1
    err = capsys.readouterr().err
    assert "unknown_sheet" in err
    assert "SENTINEL_SHEET_NAME" not in err
    assert str(workbook) not in err
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "guiskindose",
            "--settings",
            str(settings_path),
            "--file-path",
            str(_SIEMENS_DICOM),
            "--equipment-profile",
            str(json_profile),
            "--equipment-profile-sheet",
            "Profiles",
        ],
    )
    with pytest.raises(SystemExit) as excinfo:
        cli()
    assert excinfo.value.code == 1
    assert "SENTINEL" not in capsys.readouterr().err
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "guiskindose",
            "--settings",
            str(settings_path),
            "--file-path",
            str(_SIEMENS_DICOM),
            "--equipment-profile-sheet",
            "Profiles",
        ],
    )
    with pytest.raises(SystemExit) as excinfo:
        cli()
    assert excinfo.value.code == 1


@pytest.mark.parametrize("selector", [16, 10**400, "16", "9" * 5000, True, 1.5])
def test_invalid_selector_rejects_before_profile_reads(monkeypatch: pytest.MonkeyPatch, selector: Any) -> None:
    from guiskindose.cli_equipment_profiles import equipment_session_from_args
    from guiskindose.privacy import UserFacingInputError

    reads: list[bool] = []

    def reject_read(*args: Any, **kwargs: Any) -> Any:
        reads.append(True)
        raise AssertionError("unexpected profile read")

    monkeypatch.setattr("guiskindose.cli_equipment_profiles.assemble_session", reject_read)
    args = _argv_namespace(["unread-profile.xlsx"], sheet=selector)
    with pytest.raises(UserFacingInputError) as caught:
        equipment_session_from_args(args)
    assert "unread-profile" not in str(caught.value)
    assert reads == []


def test_cli_sheet_title_keeps_whitespace(tmp_path: Path) -> None:
    title = "  Profile data  "
    raw = _workbook_bytes([(title, _siemens_table())])
    path = tmp_path / "profile.xlsx"
    path.write_bytes(raw)
    args = _argv_namespace([path], sheet=title)
    actual = equipment_session_from_args(args)
    expected = assemble_session(files=[FileSource(path, sheet=0)])
    assert actual is not None
    assert actual.active_set_hash == expected.active_set_hash
