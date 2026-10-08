"""Kerma-meter factors used per tube and exam: summary fields, report tables, and exports."""

from __future__ import annotations

import io

import pandas as pd
import pytest
from calculate_dose_recursion_helpers import generate_synthetic_normalized_events
from openpyxl import load_workbook

from guiskindose import constants as c
from guiskindose import load_settings_example_json
from guiskindose.calculate_dose.calculate_dose import calculate_dose
from guiskindose.export.sections import (
    TUBE_NOTE,
    has_tube_content,
    tube_cf_text,
    tube_report_table,
    tube_source_text,
)
from guiskindose.export.writers import render_bytes
from guiskindose.export.writers.html import render_html_bytes
from guiskindose.gui.summary_formatters import format_tube_summary
from guiskindose.phantom_class import Phantom
from guiskindose.settings import PyskindoseSettings

_PLANES = {"A": "Plane A", "B": "Plane B", "single": "Single Plane"}


def _settings(**km) -> PyskindoseSettings:
    base = load_settings_example_json()
    base["mode"] = "calculate_dose"
    base["silence_pydicom_warnings"] = True
    base["phantom"]["model"] = "plane"
    base["plot"]["notebook_mode"] = False
    base["plot"]["plot_dosemap"] = False
    base["kerma_meter_correction"].update(km)
    return PyskindoseSettings(settings=base)


def _frame(tubes: list[str], stations: list[str] | None = None) -> pd.DataFrame:
    frame = generate_synthetic_normalized_events(len(tubes))
    frame["acquisition_plane"] = [_PLANES[t] for t in tubes]
    frame["acquisition_plane_canonical"] = tubes
    if stations is not None:
        frame["station_name"] = stations
        frame["device_serial"] = None  # serial would take precedence over the station
    return frame


def _summary(frame, settings, exam_id=None):
    dim = settings.phantom.dimension
    table = Phantom(phantom_model=c.PHANTOM_MODEL_TABLE, phantom_dim=dim)
    pad = Phantom(phantom_model=c.PHANTOM_MODEL_PAD, phantom_dim=dim)
    _, output, _ = calculate_dose(normalized_data=frame, settings=settings, table=table, pad=pad, exam_id=exam_id)
    assert output is not None
    return output[c.OUTPUT_KEY_TUBE_SUMMARY]


def _row(summary, tube):
    return next(r for r in summary if r["tube"] == tube)


def test_single_plane_exam_reports_its_factor_when_correction_is_on() -> None:
    summary = _summary(
        _frame(["single", "single"]), _settings(enable=True, explicit_label="u", in_memory_table={("u", "single"): 1.4})
    )
    assert has_tube_content(summary)
    row = _row(summary, "single")
    assert (row["cf_min"], row["cf_max"], row["cf_source"]) == (pytest.approx(1.4), pytest.approx(1.4), "manual")
    assert tube_cf_text(row) == "1.4"
    assert tube_source_text(row) == "manual"


def test_single_plane_exam_with_correction_off_has_no_section() -> None:
    summary = _summary(_frame(["single", "single"]), _settings())
    assert _row(summary, "single")["cf_source"] == "off"
    assert not has_tube_content(summary)


def test_biplane_with_correction_off_says_not_applied() -> None:
    summary = _summary(_frame(["A", "B"]), _settings())
    assert has_tube_content(summary)
    assert {tube_cf_text(r) for r in summary} == {"not applied"}
    assert {tube_source_text(r) for r in summary} == {"not applied"}


def test_biplane_reports_a_factor_per_tube() -> None:
    settings = _settings(enable=True, explicit_label="u", in_memory_table={("u", "A"): 1.5, ("u", "B"): 2.0})
    summary = _summary(_frame(["A", "B"]), settings)
    assert tube_cf_text(_row(summary, "A")) == "1.5"
    assert tube_cf_text(_row(summary, "B")) == "2"


def test_mixed_factors_within_a_tube_show_range_and_weighted_value() -> None:
    settings = _settings(enable=True, in_memory_table={("room-1", "A"): 1.1, ("room-2", "A"): 1.3})
    summary = _summary(_frame(["A", "A"], ["Room-1", "Room-2"]), settings)
    row = _row(summary, "A")
    assert (row["cf_min"], row["cf_max"]) == (pytest.approx(1.1), pytest.approx(1.3))
    assert 1.1 < row["applied_cf"] < 1.3
    text = tube_cf_text(row)
    assert text.startswith("1.1-1.3 (weighted ")
    assert row["cf_source"] == "manual"


def test_mixed_sources_within_a_tube_are_reported_as_mixed() -> None:
    settings = _settings(enable=True, default_factor=0.9, in_memory_table={("room-1", "A"): 1.1})
    summary = _summary(_frame(["A", "A"], ["Room-1", "Room-2"]), settings)
    assert _row(summary, "A")["cf_source"] == "mixed"


def test_file_source_is_reported(tmp_path) -> None:
    path = tmp_path / "cf.csv"
    path.write_text("equipment,tube,correction_factor\nroom-1,A,1.2\n", encoding="utf-8")
    summary = _summary(_frame(["A"], ["Room-1"]), _settings(enable=True, file=str(path)))
    assert _row(summary, "A")["cf_source"] == "file"


def test_multi_exam_table_has_a_row_per_exam_and_tube_with_different_factors() -> None:
    table = {("Exam 1", "u", "A"): 1.1, ("Exam 2", "u", "A"): 1.3, ("Exam 1", "u", "B"): 1.0, ("Exam 2", "u", "B"): 1.0}
    settings = _settings(enable=True, explicit_label="u", in_memory_table=table)
    s1 = _summary(_frame(["A", "B"]), settings, "Exam 1")
    s2 = _summary(_frame(["A", "B"]), settings, "Exam 2")
    rows = tube_report_table([("Exam 1", s1), ("Exam 2", s2)])
    assert rows[0][:2] == ["Exam", "Tube"]
    assert [(r[0], r[1], r[5]) for r in rows[1:]] == [
        ("Exam 1", "Plane A", "1.1"),
        ("Exam 1", "Plane B", "1"),
        ("Exam 2", "Plane A", "1.3"),
        ("Exam 2", "Plane B", "1"),
    ]


def test_no_equipment_labels_in_summary_tables_or_text() -> None:
    settings = _settings(enable=True, explicit_label="Secret-Room-7", in_memory_table={("secret-room-7", "A"): 1.2})
    summary = _summary(_frame(["A", "B"]), settings)
    rendered = (
        repr(summary) + repr(tube_report_table([("Exam 1", summary)])) + format_tube_summary([("Exam 1", summary)])
    )
    assert "secret" not in rendered.lower()


def test_gui_text_shows_factor_and_source() -> None:
    settings = _settings(enable=True, explicit_label="u", in_memory_table={("u", "single"): 1.4})
    text = format_tube_summary([(None, _summary(_frame(["single"]), settings))])
    assert "CF 1.4 (manual)" in text
    assert TUBE_NOTE in text


# ── exports ──────────────────────────────────────────────────────────────────


def _payload_with(summaries):
    from tests.unittests.test_export_html import _payload  # type: ignore[import-not-found]

    payload = _payload()
    payload.exams[0].tube_summary = summaries
    return payload


_ROW = {
    "tube": "single",
    "events": 2,
    "kerma_reported": 3.0,
    "kerma_corrected": 4.2,
    "applied_cf": 1.4,
    "cf_min": 1.4,
    "cf_max": 1.4,
    "cf_source": "manual",
    "peak_dose": 5.0,
}


@pytest.mark.parametrize("fmt", ["html", "xlsx"])
def test_single_plane_factor_appears_in_exports(fmt: str) -> None:
    payload = _payload_with([_ROW])
    if fmt == "html":
        html = render_html_bytes(payload).decode()
        assert "Dose by tube" in html
        assert "Factor source" in html
        assert "manual" in html
    else:
        wb = load_workbook(io.BytesIO(render_bytes(payload, "xlsx")))
        assert "Dose by tube" in wb.sheetnames
        values = [cell.value for row in wb["Dose by tube"].iter_rows() for cell in row]
        assert "manual" in values
        assert "1.4" in values


def test_docx_and_pdf_carry_the_factor_table() -> None:
    docx = pytest.importorskip("docx")
    payload = _payload_with([_ROW])
    doc = docx.Document(io.BytesIO(render_bytes(payload, "docx")))
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    assert "Factor source" in cells
    assert "manual" in cells
    pytest.importorskip("reportlab")
    assert render_bytes(payload, "pdf").startswith(b"%PDF")


def test_disabled_single_plane_exports_no_section() -> None:
    payload = _payload_with([{**_ROW, "cf_source": "off", "applied_cf": 1.0, "cf_min": 1.0, "cf_max": 1.0}])
    assert "Dose by tube" not in render_html_bytes(payload).decode()


def test_multi_exam_export_has_an_exam_column() -> None:
    from tests.unittests.test_export_xlsx import _multi_payload  # type: ignore[import-not-found]

    payload = _multi_payload()
    payload.exams[0].tube_summary = [_ROW]
    payload.exams[1].tube_summary = [{**_ROW, "applied_cf": 1.6, "cf_min": 1.6, "cf_max": 1.6}]
    html = render_html_bytes(payload).decode()
    assert "<th>Exam</th>" in html or ">Exam<" in html
    assert "1.6" in html


def test_gui_text_does_not_repeat_not_applied() -> None:
    text = format_tube_summary([(None, _summary(_frame(["A", "B"]), _settings()))])
    assert "CF not applied," in text
    assert "not applied (not applied)" not in text
