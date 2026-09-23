"""Phase 6 tests for the DOCX report writer."""

from __future__ import annotations

import io

import pandas as pd
import pytest

from guiskindose import PyskindoseSettings, load_settings_example_json
from guiskindose.export import ExportExamSource, ExportSource, collect_export_payload
from guiskindose.export.models import ImageEntry

pytest.importorskip("docx")

from guiskindose.export.writers.docx import render_docx_bytes


def _payload():
    out = {
        "psd": 1.0, "air_kerma": 3.0,
        "patient": {"patient_type": "human", "patient": {
            "human_phantom": "hudfrid",
            "patient_skin_cells": {"x": [0.0, 1.0, 2.0], "y": [0.0, 0.0, 0.0], "z": [0.0, 0.0, 0.0]},
            "triangle_vertex_indices": {"i": [0], "j": [1], "k": [2]}}, "orientation": "hfs", "offsets": {}},
        "dose_map": [[1, 1.0]],
        "corrections": {"correction_value_index": [[0, 1]], "backscatter": [[1.2, 1.4]],
                        "inverse_square_law": [[0.9, 0.8]], "medium": [1.02], "table": [0.8], "kerma": [3.0]},
    }
    s = PyskindoseSettings(settings=load_settings_example_json(), output_format="dict")
    src = ExportSource(execution_context="cli", output_dict=out,
                       exams=[ExportExamSource("e1", pd.DataFrame(), None, "e1.dcm", s, (0, 0, 0))],
                       file_name="e1.dcm")
    return collect_export_payload(src, with_images=False)


def test_write_docx_smoke():
    from docx import Document  # type: ignore[import-untyped]

    data = render_docx_bytes(_payload())
    assert data[:2] == b"PK"  # docx is a zip
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "GUISkinDose" in text


def test_write_docx_missing_image_notice():
    payload = _payload()
    payload.images = [ImageEntry(label="Dorsal", view="dorsal", exam_id=None, png_bytes=None,
                                 error_message="Image unavailable (kaleido/export error)")]
    data = render_docx_bytes(payload)  # must not raise
    assert data[:2] == b"PK"


def _handling():
    return {
        "rows": [
            {
                "event_index": 3,
                "classification": "rotational",
                "reason_codes": ["type_code_rotational", "primary_endpoint_motion"],
                "confidence": "coded",
                "requested_handling": "Auto",
                "effective_handling": "coverage",
                "fallback_reason": "",
                "ap1_start": 90.0,
                "ap2_start": 0.0,
                "ap1_end": -120.0,
                "ap2_end": 0.0,
                "primary_separation_deg": 150.0,
                "secondary_separation_deg": 0.0,
                "candidate_domain": "endpoint_paths",
                "requested_path_count": 2,
                "unique_candidate_count": 360,
                "angular_step_deg": 1.0,
                "include_static_pose": True,
                "direction_source": "unknown",
                "kerma": 0.011,
                "dap": 1.0,
                "multiplier": 1.0,
                "aggregation_rule": "max_within_sum_between",
            }
        ],
        "aggregate": {
            "total_events": 4,
            "rotational_count": 1,
            "positioner_motion_count": 0,
            "static_count": 3,
            "unknown_count": 0,
            "rotational_kerma": 0.011,
            "total_kerma": 0.05,
            "any_fallback_to_static": False,
        },
        "source_event_count": 4,
        "processed_event_count": 4,
    }


def _payload_with_handling():

    import pandas as pd

    from guiskindose import PyskindoseSettings, load_settings_example_json
    from guiskindose.export import ExportExamSource, ExportSource, collect_export_payload

    out = {
        "psd": 1.0, "air_kerma": 3.0,
        "patient": {"patient_type": "human", "patient": {
            "human_phantom": "hudfrid",
            "patient_skin_cells": {"x": [0.0, 1.0, 2.0], "y": [0.0, 0.0, 0.0], "z": [0.0, 0.0, 0.0]},
            "triangle_vertex_indices": {"i": [0], "j": [1], "k": [2]}}, "orientation": "hfs", "offsets": {}},
        "dose_map": [[1, 1.0]],
        "corrections": {"correction_value_index": [[0, 1]], "backscatter": [[1.2, 1.4]],
                        "inverse_square_law": [[0.9, 0.8]], "medium": [1.02], "table": [0.8], "kerma": [3.0]},
        "rotational_handling": _handling(),
    }
    s = PyskindoseSettings(settings=load_settings_example_json(), output_format="dict")
    src = ExportSource(execution_context="cli", output_dict=out,
                       exams=[ExportExamSource("e1", pd.DataFrame(), None, "e1.dcm", s, (0, 0, 0))],
                       file_name="e1.dcm")
    return collect_export_payload(src, with_images=False)


def test_docx_omits_rotational_section_without_handling():
    from docx import Document  # type: ignore[import-untyped]

    doc = Document(io.BytesIO(render_docx_bytes(_payload())))
    headings = [
        p.text for p in doc.paragraphs if p.style is not None and p.style.name.startswith("Heading")
    ]
    assert "Rotational handling" not in headings


def test_docx_renders_rotational_methodology_and_ledger():
    from docx import Document  # type: ignore[import-untyped]

    doc = Document(io.BytesIO(render_docx_bytes(_payload_with_handling())))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Rotational handling" in text
    assert "conditional coverage envelope" in text
    tables_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "rotational" in tables_text
    assert "coverage" in tables_text
