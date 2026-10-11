"""Unit tests for the XLSX equipment-profile adapter, its container preflight, and its worksheet XML preflight.

Workbooks are generated in memory (or under ``tmp_path``) so no tracked ``.xlsx`` binary is ever committed: a
packaged XLSX template still needs hash-pinned human clearance and is recorded as outstanding.

The regression driver for this module is two silent-loss bugs: a populated cell beyond the header was dropped
because a row conversion truncated to the header width, and a workbook that understated its declared
``<dimension>`` was silently truncated by the dimension-driven reader. Both are pinned here.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pytest

from guiskindose.equipment_profiles import (
    MAX_SOURCE_BYTES,
    EquipmentProfileError,
    ProfileErrorCode,
    active_set_hash,
    combine_documents,
    parse_csv_bytes,
    parse_json_bytes,
    parse_xlsx_bytes,
    parse_xlsx_file,
    profile_hash,
    source_bytes_hash,
    xlsx_preflight,
    xlsx_xml,
)
from guiskindose.equipment_profiles import constants as profile_constants
from guiskindose.equipment_profiles.canonical import canonical_profile_dict

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import Workbook  # noqa: E402

_EXAMPLE_DIR = (
    Path(__file__).resolve().parents[2]
    / "src"
    / "guiskindose"
    / "example_data"
    / "equipment_profiles"
)
_CSV_PATH = _EXAMPLE_DIR / "equipment_profile_example.csv"
_JSON_PATH = _EXAMPLE_DIR / "equipment_profile_example.json"

_FLAT_COLUMN_NAMES: tuple[str, ...] = tuple(
    next(csv.reader(io.StringIO(_CSV_PATH.read_text(encoding="utf-8"))))
)
_WIDTH = len(_FLAT_COLUMN_NAMES)
_PROMOTION_COLUMN = "promote_coordinate_wildcard"


def _code(excinfo: pytest.ExceptionInfo[EquipmentProfileError] | EquipmentProfileError) -> ProfileErrorCode:
    """Return an error's code, from either an ``ExceptionInfo`` or the error itself."""
    error = excinfo.value if isinstance(excinfo, pytest.ExceptionInfo) else excinfo
    return error.code


def _csv_rows() -> list[list[str]]:
    return list(csv.reader(io.StringIO(_CSV_PATH.read_text(encoding="utf-8"))))


def _typed(value: str) -> object:
    """Convert a packaged CSV cell to the typed form a workbook would store."""
    if value == "":
        return None
    if value in ("true", "false"):
        return value == "true"
    try:
        number = float(value)
    except ValueError:
        return value
    return int(number) if number == int(number) and "." not in value else number


def _workbook_bytes(
    rows: list[list[str]] | None = None,
    *,
    header: list[str] | None = None,
    mutate: Any = None,
) -> bytes:
    """Build a one-sheet workbook, typing numeric and boolean cells like a real workbook would."""
    source = _csv_rows() if rows is None else rows
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(header if header is not None else list(_FLAT_COLUMN_NAMES))
    # An explicit ``rows`` argument is used exactly as given; the default carries the CSV's own header row,
    # which is dropped here because the header is appended separately above.
    data = source[1:] if rows is None else source
    for row in data:
        sheet.append([_typed(cell) for cell in row])
    if mutate is not None:
        mutate(sheet)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _packaged_workbook_bytes() -> bytes:
    return _workbook_bytes(_csv_rows()[1:])


def _rewrite_worksheet_xml(transform: Any, *, data: bytes | None = None) -> bytes:
    """Rebuild a workbook, applying ``transform`` to each worksheet part's XML."""
    source = _packaged_workbook_bytes() if data is None else data
    target = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as archive, zipfile.ZipFile(target, "w") as writer:
        for info in archive.infolist():
            payload = archive.read(info.filename)
            if info.filename.startswith(xlsx_xml.WORKSHEET_ENTRY_PREFIX) and info.filename.endswith(
                xlsx_xml.WORKSHEET_ENTRY_SUFFIX
            ):
                payload = transform(payload.decode("utf-8")).encode("utf-8")
            writer.writestr(info, payload)
    return target.getvalue()


def _set_dimension(reference: str) -> bytes:
    return _rewrite_worksheet_xml(
        lambda text: re.sub(r'<dimension ref="[^"]*"/>', f'<dimension ref="{reference}"/>', text, count=1)
    )


def _column(name: str) -> int:
    return _FLAT_COLUMN_NAMES.index(name) + 1


# ---------------------------------------------------------------------------
# Regression: silent loss
# ---------------------------------------------------------------------------


def test_populated_cell_beyond_the_header_is_rejected_not_dropped() -> None:
    """A populated cell past the 19-column header must fail the load, never be silently discarded."""
    for value in ("=1+1", "boom", "#DIV/0!", 42, True):
        raw = _workbook_bytes(_csv_rows()[1:], mutate=lambda sheet, v=value: sheet.cell(row=2, column=20, value=v))
        with pytest.raises(EquipmentProfileError) as excinfo:
            parse_xlsx_bytes(raw)
        assert excinfo.value.source_kind == "xlsx"
        assert _code(excinfo) in (
            ProfileErrorCode.UNKNOWN_FIELD,
            ProfileErrorCode.MALFORMED_TABULAR,
            ProfileErrorCode.INVALID_TYPE,
            ProfileErrorCode.INVALID_NUMBER,
            ProfileErrorCode.UNSAFE_FORMULA,
            ProfileErrorCode.SPREADSHEET_ERROR,
        ), value


def test_blank_cell_beyond_the_header_is_tolerated_padding() -> None:
    raw = _workbook_bytes(_csv_rows()[1:], mutate=lambda sheet: sheet.cell(row=2, column=20, value=None))
    document = parse_xlsx_bytes(raw)

    assert len(document.profiles) == 3
    assert document.logical_rows == 24


def test_understated_declared_dimension_keeps_all_rows() -> None:
    """A workbook declaring far fewer rows than it holds must still be read in full.

    openpyxl's read-only parser is dimension-driven, so the adapter validates the worksheet XML and then
    discards the declared dimension. Without that, every profile after the declared extent is silently lost.
    """
    raw = _set_dimension("A1:S3")

    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        worksheet = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert '<dimension ref="A1:S3"/>' in worksheet  # the tampered metadata really is in the archive

    document = parse_xlsx_bytes(raw)
    from_csv = parse_csv_bytes(_CSV_PATH.read_bytes())
    from_json = parse_json_bytes(_JSON_PATH.read_bytes())

    assert len(document.profiles) == 3
    assert document.logical_rows == 24
    assert active_set_hash(document.profiles) == active_set_hash(from_csv.profiles)
    assert active_set_hash(document.profiles) == active_set_hash(from_json.profiles)


def test_inflated_declared_dimension_cannot_materialize_huge_padding() -> None:
    """A workbook claiming far more cells than it holds must not materialize padding from that claim.

    The declared dimension is discarded, so the reader yields the real rows. Whatever the forgery, the
    outcome is either a closed rejection or the same small document: the archive is never inflated into
    billion-cell padding, and no profile is silently lost or invented.
    """
    raw = _set_dimension("A1:XFD1048576")

    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        worksheet = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert '<dimension ref="A1:XFD1048576"/>' in worksheet

    try:
        document = parse_xlsx_bytes(raw)
    except EquipmentProfileError as excinfo:
        assert _code(excinfo) in (
            ProfileErrorCode.MALFORMED_CONTAINER,
            ProfileErrorCode.ROW_LIMIT_EXCEEDED,
        )
        return
    from_csv = parse_csv_bytes(_CSV_PATH.read_bytes())
    assert len(document.profiles) == len(from_csv.profiles)
    assert active_set_hash(document.profiles) == active_set_hash(from_csv.profiles)


def test_untampered_workbook_matches_csv_and_json_canonically() -> None:
    document = parse_xlsx_bytes(_packaged_workbook_bytes())
    from_csv = parse_csv_bytes(_CSV_PATH.read_bytes())
    from_json = parse_json_bytes(_JSON_PATH.read_bytes())

    assert [p.profile_id for p in document.profiles] == [p.profile_id for p in from_csv.profiles]
    for left, right in zip(from_json.profiles, document.profiles, strict=True):
        assert canonical_profile_dict(left) == canonical_profile_dict(right)
        assert profile_hash(left) == profile_hash(right)
    assert active_set_hash(from_json.profiles) == active_set_hash(document.profiles)
    assert active_set_hash(from_csv.profiles) == active_set_hash(document.profiles)
    assert document.source_hash == source_bytes_hash(_packaged_workbook_bytes())


# ---------------------------------------------------------------------------
# Worksheet XML preflight
# ---------------------------------------------------------------------------


def test_worksheet_xml_bounds_are_actual_coordinates() -> None:
    xml = (
        b'<?xml version="1.0"?><worksheet><sheetData>'
        b'<row r="1"><c r="A1" t="inlineStr"><is><t>a</t></is></c><c r="B1"/></row>'
        b'<row r="3"><c r="A3"/></row>'
        b"</sheetData></worksheet>"
    )
    assert xlsx_xml.validate_worksheet_xml(xml, source_kind="xlsx") == (2, 2, 3)


def test_worksheet_xml_rejects_a_declaration_of_doctype() -> None:
    payload = b'<?xml version="1.0"?><!DOCTYPE sheet [<!ENTITY x "y">]><sheetData/>'
    with pytest.raises(EquipmentProfileError) as excinfo:
        xlsx_xml.validate_worksheet_xml(payload, source_kind="xlsx")
    assert _code(excinfo) is ProfileErrorCode.MALFORMED_CONTAINER


def test_worksheet_xml_rejects_truncated_and_non_xml_payloads() -> None:
    good = (
        b'<?xml version="1.0"?><worksheet><sheetData><row r="1"><c r="A1"/></row></sheetData></worksheet>'
    )
    for payload in (good[: len(good) // 2], b"not xml at all", b"", b'<?xml version="1.0"?><a></b>'):
        with pytest.raises(EquipmentProfileError) as excinfo:
            xlsx_xml.validate_worksheet_xml(payload, source_kind="xlsx")
        assert _code(excinfo) is ProfileErrorCode.MALFORMED_CONTAINER
        assert excinfo.value.__cause__ is None


def test_worksheet_xml_rejects_repeated_and_malformed_coordinates() -> None:
    repeated = (
        b'<?xml version="1.0"?><worksheet><sheetData>'
        b'<row r="1"><c r="A1"/><c r="A1"/></row>'
        b"</sheetData></worksheet>"
    )
    with pytest.raises(EquipmentProfileError) as excinfo:
        xlsx_xml.validate_worksheet_xml(repeated, source_kind="xlsx")
    assert _code(excinfo) is ProfileErrorCode.MALFORMED_CONTAINER

    duplicated_row = (
        b'<?xml version="1.0"?><worksheet><sheetData>'
        b'<row r="1"><c r="A1"/></row><row r="1"><c r="B1"/></row>'
        b"</sheetData></worksheet>"
    )
    with pytest.raises(EquipmentProfileError) as excinfo:
        xlsx_xml.validate_worksheet_xml(duplicated_row, source_kind="xlsx")
    assert _code(excinfo) is ProfileErrorCode.MALFORMED_CONTAINER


def test_worksheet_xml_rejects_out_of_range_coordinates() -> None:
    for column, row in (("A", 0), ("A", 1_048_577), ("ZZZZZZ", 1), ("A1", None)):
        reference = f"{column}{row}" if row is not None else "A0"
        payload = (
            b'<?xml version="1.0"?><worksheet><sheetData>'
            + f'<row r="{row}"><c r="{reference}"/></row></sheetData></worksheet>'.encode()
            if row is not None
            else b'<?xml version="1.0"?><worksheet><sheetData><c r="A0"/></sheetData></worksheet>'
        )
        with pytest.raises(EquipmentProfileError) as excinfo:
            xlsx_xml.validate_worksheet_xml(payload, source_kind="xlsx")
        assert _code(excinfo) is ProfileErrorCode.MALFORMED_CONTAINER, (column, row)


def test_worksheet_xml_column_reference_helpers() -> None:
    assert xlsx_xml.column_index("A") == 1
    assert xlsx_xml.column_index("Z") == 26
    assert xlsx_xml.column_index("AA") == 27
    assert xlsx_xml.column_index("XFD") == 16_384
    assert xlsx_xml.column_index("a1") is None
    assert xlsx_xml.cell_reference("A0") is None
    assert xlsx_xml.cell_reference("A٣") is None
    assert xlsx_xml.cell_reference("A1٣") is None
    assert xlsx_xml.cell_reference("ZZZZ") is None
    assert xlsx_xml.cell_reference("XFD1048576") == (1_048_576, 16_384)


def test_actual_row_budget_is_enforced_from_worksheet_xml(monkeypatch: pytest.MonkeyPatch) -> None:
    """A sheet whose real row count exceeds the budget fails before openpyxl materializes it."""
    monkeypatch.setattr(profile_constants, "MAX_XLSX_PREFLIGHT_ROWS", 5)
    raw = _workbook_bytes(_csv_rows()[1:])

    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(raw)
    assert _code(excinfo) is ProfileErrorCode.MALFORMED_CONTAINER


def test_actual_column_budget_is_enforced_from_worksheet_xml(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(profile_constants, "MAX_XLSX_PREFLIGHT_COLS", 3)
    workbook_bytes = _packaged_workbook_bytes()
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.MALFORMED_CONTAINER


def test_actual_cell_budget_is_enforced_from_worksheet_xml(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(profile_constants, "MAX_XLSX_PREFLIGHT_CELLS", 20)
    workbook_bytes = _packaged_workbook_bytes()
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.MALFORMED_CONTAINER


def test_corrupt_worksheet_xml_with_a_valid_crc_fails_closed() -> None:
    """A corrupt selected worksheet with an intact archive fails closed, with no library cause leaked."""
    raw = _rewrite_worksheet_xml(
        lambda text: '<?xml version="1.0"?><worksheet><sheetData><row r="1"><c r="A1"><v>'
    )

    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(raw)

    assert _code(excinfo) is ProfileErrorCode.MALFORMED_CONTAINER
    assert excinfo.value.__cause__ is None
    assert excinfo.value.source_kind == "xlsx"


# ---------------------------------------------------------------------------
# Sheet selection
# ---------------------------------------------------------------------------


def _multi_sheet_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Profiles"
    first.append(list(_FLAT_COLUMN_NAMES))
    for row in _csv_rows()[1:]:
        first.append([_typed(cell) for cell in row])
    second = workbook.create_sheet("Other")
    second.append(["not", "a", "profile"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_multi_sheet_workbook_defaults_to_first_sheet() -> None:
    default = parse_xlsx_bytes(_multi_sheet_bytes())
    explicit = parse_xlsx_bytes(_multi_sheet_bytes(), sheet=0)
    assert active_set_hash(default.profiles) == active_set_hash(explicit.profiles)


def test_selected_non_profile_sheet_fails_validation() -> None:
    workbook_bytes = _multi_sheet_bytes()
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes, sheet=1)
    assert _code(excinfo) is ProfileErrorCode.UNKNOWN_FIELD


def test_out_of_range_sheet_selection_is_unknown_sheet() -> None:
    workbook_bytes = _multi_sheet_bytes()
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes, sheet=9)
    assert _code(excinfo) is ProfileErrorCode.UNKNOWN_SHEET

    workbook_bytes = _multi_sheet_bytes()
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes, sheet=-1)
    assert _code(excinfo) is ProfileErrorCode.UNKNOWN_SHEET


def test_sheet_count_comes_from_the_workbook_catalog(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(profile_constants, "MAX_XLSX_SHEETS", 1)
    workbook_bytes = _multi_sheet_bytes()
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes, sheet=0)
    assert _code(excinfo) is ProfileErrorCode.SHEET_LIMIT_EXCEEDED


# ---------------------------------------------------------------------------
# Typed cell policy
# ---------------------------------------------------------------------------


def _wildcard_rows() -> list[list[str]]:
    """Data rows for one manufacturer-scoped wildcard profile with a single coordinate leaf."""
    return [
        ["profile", "demo-wild", "Acme Imaging", "*", "", "", "s", "n", *([""] * 11)],
        ["coordinate", "demo-wild", *([""] * 6), "translation_offset.z", "-1.0", *([""] * 9)],
    ]


def _tx_row() -> int:
    """The 1-based sheet row of the first packaged transmission record."""
    for index, row in enumerate(_csv_rows(), start=1):
        if row[0] == "transmission":
            return index
    raise AssertionError("packaged example has no transmission row")


def _mutate(row: int, column: str, value: object) -> Any:
    def mutate(sheet: Any) -> None:
        sheet.cell(row=row, column=_column(column), value=value)

    return mutate


def test_boolean_in_identity_and_free_text_is_a_type_error() -> None:
    for column in ("manufacturer", "model_aliases", "manufacturer_aliases", "notes", "source"):
        workbook_bytes = _workbook_bytes(_csv_rows()[1:], mutate=_mutate(2, column, True))
        with pytest.raises(EquipmentProfileError) as excinfo:
            parse_xlsx_bytes(workbook_bytes)
        assert _code(excinfo) is ProfileErrorCode.INVALID_TYPE, column


def test_number_in_identity_and_header_is_a_type_error() -> None:
    for column in ("manufacturer", "record_type", "models"):
        workbook_bytes = _workbook_bytes(_csv_rows()[1:], mutate=_mutate(2, column, 12345))
        with pytest.raises(EquipmentProfileError) as excinfo:
            parse_xlsx_bytes(workbook_bytes)
        assert _code(excinfo) is ProfileErrorCode.INVALID_TYPE, column

    header: list[Any] = list(_FLAT_COLUMN_NAMES)
    header[4] = 42
    workbook_bytes = _workbook_bytes(_csv_rows()[1:], header=header)
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.INVALID_TYPE


def test_boolean_in_a_numeric_column_is_an_invalid_number() -> None:
    for column in ("kvp_kv", "filtration_added_mmcu", "filtration_added_mmal", "k_patient_support"):
        workbook_bytes = _workbook_bytes(_csv_rows()[1:], mutate=_mutate(_tx_row(), column, True))
        with pytest.raises(EquipmentProfileError) as excinfo:
            parse_xlsx_bytes(workbook_bytes)
        assert _code(excinfo) is ProfileErrorCode.INVALID_NUMBER, column


def test_boolean_is_allowed_only_in_a_boolean_context() -> None:
    """The swap leaf and the promotion column are the only boolean contexts."""
    swap_row = next(
        index for index, row in enumerate(_csv_rows(), start=1)
        if row[0] == "coordinate" and row[8] == "swap_lateral_longitudinal"
    )
    document = parse_xlsx_bytes(_workbook_bytes(_csv_rows()[1:], mutate=_mutate(swap_row, "value", True)))
    coordinates = document.profiles[0].coordinates
    assert coordinates is not None
    assert coordinates.swap_lateral_longitudinal is True

    promoted_header = [*_FLAT_COLUMN_NAMES, _PROMOTION_COLUMN]

    def promote(sheet: Any) -> None:
        sheet.cell(row=2, column=len(_FLAT_COLUMN_NAMES) + 1, value=True)

    document = parse_xlsx_bytes(
        _workbook_bytes(_wildcard_rows(), header=promoted_header, mutate=promote)
    )
    assert document.profiles[0].promote_coordinate_wildcard is True


def test_boolean_in_a_sign_or_mode_leaf_is_a_type_error() -> None:
    sign_row = next(
        index for index, row in enumerate(_csv_rows(), start=1)
        if row[0] == "coordinate" and row[8] == "translation_direction.x"
    )
    workbook_bytes = _workbook_bytes(_csv_rows()[1:], mutate=_mutate(sign_row, "value", True))
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.INVALID_TYPE

    mode_row = next(
        index for index, row in enumerate(_csv_rows(), start=1)
        if row[0] == "coordinate" and row[8] == "field_size_mode"
    )
    workbook_bytes = _workbook_bytes(_csv_rows()[1:], mutate=_mutate(mode_row, "value", 7))
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.INVALID_TYPE

    workbook_bytes = _workbook_bytes(_csv_rows()[1:], mutate=_mutate(sign_row, "value", 1))
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.INVALID_TYPE


def test_a_literal_boolean_string_is_allowed_where_text_is_allowed() -> None:
    document = parse_xlsx_bytes(_workbook_bytes(_csv_rows()[1:], mutate=_mutate(2, "manufacturer", "true")))
    assert document.profiles[0].manufacturer == "true"


def test_dates_are_valid_only_in_the_date_column() -> None:
    """A typed date in the date column renders strict ISO; anywhere else it is a type error."""
    document = parse_xlsx_bytes(
        _workbook_bytes(_csv_rows()[1:], mutate=_mutate(_tx_row(), "date", date(1901, 1, 1)))
    )
    assert document.profiles[0].transmission[0].date == "1901-01-01"

    workbook_bytes = _workbook_bytes(_csv_rows()[1:], mutate=_mutate(2, "models", date(2026, 1, 1)))
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) in (
        ProfileErrorCode.INVALID_DATE_OBJECT,
        ProfileErrorCode.INVALID_TYPE,
    )

    workbook_bytes = _workbook_bytes(_csv_rows()[1:], mutate=_mutate(_tx_row(), "date", datetime(1901, 1, 1, 12, 0)))
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(
            workbook_bytes
        )
    assert _code(excinfo) is ProfileErrorCode.INVALID_DATE_OBJECT


def test_formula_and_error_cells_are_rejected_everywhere() -> None:
    for column in ("models", "value", "date", "notes"):
        workbook_bytes = _workbook_bytes(_csv_rows()[1:], mutate=_mutate(2, column, "=1+1"))
        with pytest.raises(EquipmentProfileError) as excinfo:
            parse_xlsx_bytes(workbook_bytes)
        assert _code(excinfo) is ProfileErrorCode.UNSAFE_FORMULA, column

    workbook_bytes = _workbook_bytes(_csv_rows()[1:], mutate=_mutate(2, "models", "#DIV/0!"))
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.SPREADSHEET_ERROR


def _blank_cell(column: str) -> list[list[str]]:
    """Return the packaged data rows with one transmission column blanked."""
    position = _FLAT_COLUMN_NAMES.index(column)
    rows = [list(row) for row in _csv_rows()[1:]]
    target = next(index for index, row in enumerate(rows) if row[0] == "transmission")
    rows[target][position] = ""
    return rows


def test_blank_filtration_is_rejected() -> None:
    workbook_bytes = _workbook_bytes(_blank_cell("filtration_added_mmcu"))
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.BLANK_DISALLOWED_CELL

    workbook_bytes = _workbook_bytes(_blank_cell("filtration_added_mmal"))
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.BLANK_DISALLOWED_CELL


def test_explicit_zero_filtration_is_accepted() -> None:
    document = parse_xlsx_bytes(
        _workbook_bytes(
            _blank_cell("filtration_added_mmcu"),  # blank in text; the typed writer turns "" into None
            mutate=lambda sheet: sheet.cell(
                row=_tx_row(), column=_column("filtration_added_mmcu"), value=0
            ),
        )
    )
    assert document.profiles[0].transmission[0].filtration_added_mmcu == 0.0


def test_duplicate_coordinate_leaf_and_orphan_rules_are_reused() -> None:
    leaf_rows = [row for row in _csv_rows()[1:] if row[0] == "coordinate"]
    profile_rows = [row for row in _csv_rows()[1:] if row[0] == "profile"]
    duplicate = [*profile_rows, *leaf_rows, leaf_rows[0]]

    workbook_bytes = _workbook_bytes(duplicate)
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.DUPLICATE_JSON_KEY

    workbook_bytes = _workbook_bytes(leaf_rows)
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.MISSING_REQUIRED_FIELD


def test_duplicate_transmission_key_fails_across_documents() -> None:
    rows = _csv_rows()[1:]
    first = parse_xlsx_bytes(_workbook_bytes(rows))
    second = parse_xlsx_bytes(_workbook_bytes(rows))
    with pytest.raises(EquipmentProfileError) as excinfo:
        combine_documents([first, second])
    assert _code(excinfo) is ProfileErrorCode.DUPLICATE_PROFILE_ID


def test_twenty_column_header_accepts_promotion_on_a_wildcard() -> None:
    """The optional promotion column is a valid 20th column and promotion needs a wildcard model."""
    header = [*_FLAT_COLUMN_NAMES, _PROMOTION_COLUMN]

    def promote(sheet: Any) -> None:
        sheet.cell(row=2, column=len(_FLAT_COLUMN_NAMES) + 1, value=True)

    document = parse_xlsx_bytes(_workbook_bytes(_wildcard_rows(), header=header, mutate=promote))
    assert document.profiles[0].promote_coordinate_wildcard is True


def test_promotion_on_a_non_wildcard_profile_is_rejected() -> None:
    header = [*_FLAT_COLUMN_NAMES, _PROMOTION_COLUMN]

    def promote(sheet: Any) -> None:
        sheet.cell(row=4, column=len(_FLAT_COLUMN_NAMES) + 1, value=True)

    rows = [*_wildcard_rows(), *_csv_rows()[1:]]
    workbook_bytes = _workbook_bytes(rows, header=header, mutate=promote)
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.WILDCARD_NOT_PERMITTED


# ---------------------------------------------------------------------------
# Container preflight
# ---------------------------------------------------------------------------


def test_non_zip_bytes_fail_closed() -> None:
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(b"this is not a zip archive")
    assert _code(excinfo) is ProfileErrorCode.MALFORMED_CONTAINER
    assert excinfo.value.__cause__ is None


def test_duplicate_archive_entry_names_are_rejected() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/worksheets/sheet1.xml", "<sheet/>")
        archive.writestr("xl/worksheets/sheet1.xml", "<sheet/>")

    workbook_bytes = buffer.getvalue()
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.MALFORMED_CONTAINER


def test_inflation_lying_metadata_is_caught(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(profile_constants, "MAX_XLSX_MEMBER_BYTES", 64)
    workbook_bytes = _packaged_workbook_bytes()
    with pytest.raises(EquipmentProfileError) as excinfo:
        xlsx_preflight.preflight_container(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.INFLATION_LIMIT_EXCEEDED

    monkeypatch.setattr(profile_constants, "MAX_XLSX_MEMBER_BYTES", 1024 * 1024)
    monkeypatch.setattr(profile_constants, "MAX_XLSX_TOTAL_INFLATED_BYTES", 200)
    workbook_bytes = _packaged_workbook_bytes()
    with pytest.raises(EquipmentProfileError) as excinfo:
        xlsx_preflight.preflight_container(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.INFLATION_LIMIT_EXCEEDED


def test_truthful_inflation_passes() -> None:
    monkeypatch_limits = {
        "MAX_XLSX_MEMBER_BYTES": 1024 * 1024,
        "MAX_XLSX_TOTAL_INFLATED_BYTES": 8 * 1024 * 1024,
    }
    original = {name: getattr(profile_constants, name) for name in monkeypatch_limits}
    for name, value in monkeypatch_limits.items():
        setattr(profile_constants, name, value)
    try:
        assert xlsx_preflight.preflight_container(_packaged_workbook_bytes()) >= 1
        assert parse_xlsx_bytes(_packaged_workbook_bytes()).source_kind == "xlsx"
    finally:
        for name, value in original.items():
            setattr(profile_constants, name, value)


def test_oversized_source_is_rejected_by_the_bounded_read(tmp_path: Path) -> None:
    target = tmp_path / "big.xlsx"
    target.write_bytes(b"x" * (MAX_SOURCE_BYTES + 1))
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_file(target)
    assert _code(excinfo) is ProfileErrorCode.OVERSIZED_SOURCE
    assert str(target) not in str(excinfo.value)


def test_unreadable_file_and_bad_suffix_use_the_xlsx_source_kind(tmp_path: Path) -> None:
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_file(tmp_path / "missing.xlsx")
    assert excinfo.value.source_kind == "xlsx"
    assert _code(excinfo) is ProfileErrorCode.UNREADABLE_SOURCE

    wrong = tmp_path / "profiles.csv"
    wrong.write_bytes(b"x")
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_file(wrong)
    assert _code(excinfo) is ProfileErrorCode.UNSUPPORTED_SOURCE_SUFFIX
    assert excinfo.value.source_kind == "xlsx"


def test_file_adapter_matches_bytes_adapter(tmp_path: Path) -> None:
    raw = _packaged_workbook_bytes()
    target = tmp_path / "profiles.xlsx"
    target.write_bytes(raw)

    from_file = parse_xlsx_file(target)
    from_bytes = parse_xlsx_bytes(raw)

    assert from_file.source_hash == from_bytes.source_hash
    assert [profile_hash(p) for p in from_file.profiles] == [profile_hash(p) for p in from_bytes.profiles]


def test_workbook_close_failure_does_not_mask_a_parse(monkeypatch: pytest.MonkeyPatch) -> None:
    """A workbook whose close() raises still yields its parsed profiles."""
    def _boom() -> None:
        raise RuntimeError("close failed")

    monkeypatch.setattr(Workbook, "close", _boom, raising=True)
    document = parse_xlsx_bytes(_packaged_workbook_bytes())
    assert len(document.profiles) == 3


def test_no_error_echoes_sheet_names_paths_or_cells() -> None:
    workbook_bytes = _multi_sheet_bytes()
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes, sheet=9)
    text = str(excinfo.value) + repr(excinfo.value)
    assert "Profiles" not in text
    assert "Other" not in text
    assert "demo-partial" not in text
    assert "Acme" not in text

    workbook_bytes = _workbook_bytes(_csv_rows()[1:], mutate=_mutate(2, "notes", "=SECRET-NOTE"))
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert _code(excinfo) is ProfileErrorCode.UNSAFE_FORMULA
    assert "SECRET-NOTE" not in str(excinfo.value)
    assert "SECRET-NOTE" not in repr(excinfo.value)


# ---------------------------------------------------------------------------
# Laziness of the third-party dependency
# ---------------------------------------------------------------------------


def test_openpyxl_is_not_required_for_json_or_csv_parsing() -> None:
    """The adapter must not pull the spreadsheet stack in for a JSON or CSV parse.

    The top-level ``guiskindose`` package imports openpyxl through the existing tabular loader, so this drops
    it from ``sys.modules`` first and then verifies a JSON/CSV parse does not re-add it.
    """
    import sys

    from guiskindose.equipment_profiles import csv_source, json_source

    for name in [name for name in sys.modules if name == "openpyxl" or name.startswith("openpyxl.")]:
        del sys.modules[name]

    assert parse_json_bytes(_JSON_PATH.read_bytes()).profiles
    assert parse_csv_bytes(_CSV_PATH.read_bytes()).profiles
    assert "openpyxl" not in sys.modules
    assert csv_source is not None
    assert json_source is not None


def test_canonical_import_is_reachable() -> None:
    from guiskindose.equipment_profiles.canonical import canonical_profile_dict

    document = parse_xlsx_bytes(_packaged_workbook_bytes())
    assert canonical_profile_dict(document.profiles[0])["profile_id"] == "demo-scope-2000"


@pytest.mark.parametrize("body", [
    '<row r="999999"><c r="A999999"/></row>',
    '<row r="1"><c r="A2"/></row>',
    '<row r="2"><c r="A2"/></row><row r="1"><c r="A1"/></row>',
    '<row r="1"><c r="B1"/><c r="A1"/></row>',
])
def test_xml_coordinate_order_and_extent_are_checked(body: str) -> None:
    xml = ('<worksheet><sheetData>' + body + '</sheetData></worksheet>').encode()
    with pytest.raises(EquipmentProfileError):
        xlsx_xml.validate_worksheet_xml(xml, source_kind="xlsx")


def test_namespace_prefix_does_not_bypass_xml_bounds() -> None:
    xml = b'<s:worksheet xmlns:s="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><s:sheetData><s:row r="1"><s:c r="ZZ1"/></s:row></s:sheetData></s:worksheet>'
    with pytest.raises(EquipmentProfileError):
        xlsx_xml.validate_worksheet_xml(xml, source_kind="xlsx")


def test_prefixed_worksheet_retains_all_profiles_and_hashes() -> None:
    from xml.etree import ElementTree

    raw = _rewrite_worksheet_xml(lambda text: ElementTree.tostring(ElementTree.fromstring(text)).decode())
    parsed = parse_xlsx_bytes(raw)
    expected = parse_csv_bytes(_CSV_PATH.read_bytes())
    assert active_set_hash(parsed.profiles) == active_set_hash(expected.profiles)


def test_text_cells_use_the_same_numeric_and_boolean_grammar_as_csv() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in _csv_rows():
        sheet.append(row)
    raw = io.BytesIO()
    workbook.save(raw)
    parsed = parse_xlsx_bytes(raw.getvalue())
    expected = parse_csv_bytes(_CSV_PATH.read_bytes())
    assert active_set_hash(parsed.profiles) == active_set_hash(expected.profiles)


def test_nonempty_whitespace_extra_column_is_not_discarded() -> None:
    def add_extra(sheet: Any) -> None:
        sheet.cell(row=2, column=20, value=" ")
    workbook_bytes = _workbook_bytes(mutate=add_extra)
    with pytest.raises(EquipmentProfileError) as excinfo:
        parse_xlsx_bytes(workbook_bytes)
    assert excinfo.value.code == ProfileErrorCode.UNKNOWN_FIELD


@pytest.mark.parametrize("invalid", [False, True])
def test_nonstandard_worksheet_part_is_still_preflighted(invalid: bool) -> None:
    source = _packaged_workbook_bytes()
    output = io.BytesIO()
    old = "xl/worksheets/sheet1.xml"
    new = "xl/custom/profile-data.xml"
    with zipfile.ZipFile(io.BytesIO(source)) as archive, zipfile.ZipFile(output, "w") as target:
        for item in archive.infolist():
            payload = archive.read(item)
            if item.filename in {"[Content_Types].xml", "xl/_rels/workbook.xml.rels"}:
                payload = payload.replace(old.encode(), new.encode())
            if item.filename == old and invalid:
                payload = payload.replace(b'<row r="2">', b'<row r="999999">')
            target.writestr(new if item.filename == old else item.filename, payload)
    if invalid:
        workbook_bytes = output.getvalue()
        with pytest.raises(EquipmentProfileError):
            parse_xlsx_bytes(workbook_bytes)
    else:
        actual = parse_xlsx_bytes(output.getvalue())
        expected = parse_csv_bytes(_CSV_PATH.read_bytes())
        assert active_set_hash(actual.profiles) == active_set_hash(expected.profiles)
