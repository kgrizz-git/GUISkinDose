"""Equipment-profile match formatting and export table helpers."""

from __future__ import annotations

import pandas as pd

from guiskindose import PyskindoseSettings, load_settings_example_json
from guiskindose.equipment_profile_display import (
    EQUIPMENT_PROFILE_SECTION_TITLE,
    equipment_match_kind_label,
    equipment_profile_report_table,
    format_equipment_profile_blocks,
    has_equipment_profile_content,
)
from guiskindose.export import ExportExamSource, ExportSource, collect_export_payload
from guiskindose.export.sections import equipment_profile_blocks
from guiskindose.export.writers.html import render_html_bytes


def _sample_match(kind: str = "partial") -> dict:
    return {
        "profile_id": "p-test",
        "source": "custom",
        "level": "exact",
        "notice": {
            "kind": kind,
            "entries": [],
            "text": "partial match — fields inherited from bundled/Default",
        },
    }


def test_has_equipment_profile_content_requires_notice_text() -> None:
    assert has_equipment_profile_content(_sample_match())
    assert not has_equipment_profile_content(None)
    assert not has_equipment_profile_content({"profile_id": "x", "notice": {}})


def test_equipment_match_kind_label_unknown() -> None:
    assert equipment_match_kind_label(_sample_match("weird")) == "Unknown"
    assert equipment_match_kind_label({"profile_id": "p", "notice": {}}) == "Unknown"


def test_format_and_table_rows() -> None:
    match = _sample_match("complete")
    assert equipment_match_kind_label(match) == "Complete"
    text = format_equipment_profile_blocks([(None, match)])
    assert "p-test" in text
    assert "Complete" in text
    rows = equipment_profile_report_table([(None, match)])
    assert rows[0] == ["Profile", "Source", "Match", "Notice"]
    assert rows[1][0] == "p-test"


def test_equipment_profile_blocks_on_payload() -> None:
    from types import SimpleNamespace

    match = _sample_match()
    exam = SimpleNamespace(exam_id="Exam 1", equipment_profile_match=match)
    payload = SimpleNamespace(is_multi_exam=False, exams=[exam])
    assert equipment_profile_blocks(payload) == [(None, match)]


def test_equipment_profile_blocks_multi_exam_labels() -> None:
    from types import SimpleNamespace

    m1, m2 = _sample_match(), _sample_match()
    payload = SimpleNamespace(
        is_multi_exam=True,
        exams=[
            SimpleNamespace(exam_id="Study A", equipment_profile_match=m1),
            SimpleNamespace(exam_id="Study B", equipment_profile_match=m2),
        ],
    )
    assert equipment_profile_blocks(payload) == [("Study A", m1), ("Study B", m2)]
    rows = equipment_profile_report_table(equipment_profile_blocks(payload))
    assert rows[0][0] == "Exam"
    assert rows[1][0] == "Study A"


def test_html_export_includes_equipment_profile_section() -> None:
    out = {
        "psd": 1.0,
        "air_kerma": 3.0,
        "equipment_profile_match": _sample_match(),
        "patient": {
            "patient_type": "human",
            "patient": {
                "human_phantom": "hudfrid",
                "patient_skin_cells": {"x": [0.0], "y": [0.0], "z": [0.0]},
                "triangle_vertex_indices": {"i": [0], "j": [0], "k": [0]},
            },
            "orientation": "hfs",
            "offsets": {},
        },
        "dose_map": [[0, 1.0]],
        "corrections": {
            "correction_value_index": [[0]],
            "backscatter": [[1.0]],
            "inverse_square_law": [[1.0]],
            "medium": [1.0],
            "table": [1.0],
            "kerma": [3.0],
        },
    }
    settings = PyskindoseSettings(settings=load_settings_example_json(), output_format="dict")
    source = ExportSource(
        execution_context="cli",
        output_dict=out,
        exams=[ExportExamSource("e1", pd.DataFrame(), None, "e1.dcm", settings, (0, 0, 0))],
        file_name="e1.dcm",
    )
    payload = collect_export_payload(source, with_images=False)
    html = render_html_bytes(payload).decode()
    assert EQUIPMENT_PROFILE_SECTION_TITLE in html
    assert "p-test" in html


def _export_payload(with_match: bool):
    out = {
        "psd": 1.0,
        "air_kerma": 3.0,
        "patient": {
            "patient_type": "human",
            "patient": {
                "human_phantom": "hudfrid",
                "patient_skin_cells": {"x": [0.0], "y": [0.0], "z": [0.0]},
                "triangle_vertex_indices": {"i": [0], "j": [0], "k": [0]},
            },
            "orientation": "hfs",
            "offsets": {},
        },
        "dose_map": [[0, 1.0]],
        "corrections": {
            "correction_value_index": [[0]],
            "backscatter": [[1.0]],
            "inverse_square_law": [[1.0]],
            "medium": [1.0],
            "table": [1.0],
            "kerma": [3.0],
        },
    }
    if with_match:
        out["equipment_profile_match"] = _sample_match()
    settings = PyskindoseSettings(settings=load_settings_example_json(), output_format="dict")
    source = ExportSource(
        execution_context="cli",
        output_dict=out,
        exams=[ExportExamSource("e1", pd.DataFrame(), None, "e1.dcm", settings, (0, 0, 0))],
        file_name="e1.dcm",
    )
    return collect_export_payload(source, with_images=False)


def _docx_body_text(docx_bytes: bytes) -> str:
    import io
    import zipfile

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
        return archive.read("word/document.xml").decode()


def test_exports_omit_equipment_section_without_match() -> None:
    import io

    import pytest
    from openpyxl import load_workbook

    from guiskindose.export.writers import render_bytes

    payload = _export_payload(with_match=False)
    html = render_html_bytes(payload).decode()
    assert EQUIPMENT_PROFILE_SECTION_TITLE not in html
    xlsx = render_bytes(payload, "xlsx")
    wb = load_workbook(io.BytesIO(xlsx))
    assert "Equipment profile" not in wb.sheetnames
    assert EQUIPMENT_PROFILE_SECTION_TITLE not in _docx_body_text(render_bytes(payload, "docx"))
    pytest.importorskip("reportlab")
    from guiskindose.export.writers.pdf import _equipment_profile_flow

    assert _equipment_profile_flow(payload) == []
    assert render_bytes(payload, "pdf").startswith(b"%PDF")


def test_exports_include_equipment_section_when_match_present() -> None:
    import io

    import pytest
    from openpyxl import load_workbook

    from guiskindose.export.writers import render_bytes

    payload = _export_payload(with_match=True)
    html = render_html_bytes(payload).decode()
    assert EQUIPMENT_PROFILE_SECTION_TITLE in html
    xlsx = render_bytes(payload, "xlsx")
    wb = load_workbook(io.BytesIO(xlsx))
    assert "Equipment profile" in wb.sheetnames
    assert EQUIPMENT_PROFILE_SECTION_TITLE in _docx_body_text(render_bytes(payload, "docx"))
    pytest.importorskip("reportlab")
    from guiskindose.export.writers.pdf import _equipment_profile_flow

    assert _equipment_profile_flow(payload)
    assert render_bytes(payload, "pdf").startswith(b"%PDF")
