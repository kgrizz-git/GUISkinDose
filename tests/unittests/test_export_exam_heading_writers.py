"""Rendered reports must print each exam id once: ``Exam 1``, never ``Exam Exam 1``."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from guiskindose.export.writers.html import render_html_bytes
from guiskindose.export.writers.xlsx import render_xlsx_bytes
from tests.unittests.test_export_xlsx import _multi_payload  # type: ignore[import-not-found]


def test_html_report_prints_exam_ids_once() -> None:
    html = render_html_bytes(_multi_payload()).decode("utf-8")
    assert "<summary>Exam 1</summary>" in html
    assert "Exam Exam" not in html


def test_xlsx_report_prints_exam_ids_once() -> None:
    workbook = load_workbook(io.BytesIO(render_xlsx_bytes(_multi_payload())))
    cells = [str(c.value) for ws in workbook.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("Exam 1" in cell for cell in cells)
    assert not any("Exam Exam" in cell for cell in cells)


def test_docx_report_prints_exam_ids_once() -> None:
    pytest.importorskip("docx")
    from guiskindose.export.writers.docx import build_document

    document = build_document(_multi_payload())
    texts = [p.text for p in document.paragraphs]
    assert not any("Exam Exam" in text for text in texts)


def test_pdf_report_prints_exam_ids_once() -> None:
    pytest.importorskip("reportlab")
    from reportlab.platypus import Paragraph

    from guiskindose.export.writers.pdf import _story

    texts = [flow.text for flow in _story(_multi_payload()) if isinstance(flow, Paragraph)]
    assert not any("Exam Exam" in text for text in texts)
