"""XLSX equipment-profile adapter.

Reads a bounded byte snapshot once, preflights the OOXML container and its worksheet XML, then reads the selected
sheet and assembles the same canonical document mapping the JSON, CSV, and TSV adapters produce — so a workbook
and the equivalent CSV have identical profile hashes while their source hashes differ.

Design constraints this module exists to hold:

- **The canonical core stays stdlib-only.** openpyxl is imported inside the functions that need it, so a caller
  that only parses JSON, CSV, or TSV never imports the spreadsheet stack.
- **Nothing is trusted from metadata.** Archive entry lists, declared member sizes, and declared worksheet
  dimensions are all checked against real bytes and real coordinates. Because openpyxl's read-only parser is
  dimension-driven, the adapter validates the worksheet XML itself and then calls ``reset_dimensions()`` so a
  workbook that understates its content can never be silently truncated.
- **No round trip through CSV.** Typed worksheet cells are converted straight to the flat text cells the shared
  :func:`guiskindose.equipment_profiles.csv_source.rows_to_document` seam consumes, so there is exactly one flat
  parsing and validation implementation and no serialization loss.
- **Typed cell policy, from the schema.** Numeric cells stay numbers, and a boolean-typed cell is *not* a
  number. Booleans are accepted only where the schema expects a boolean; a boolean in an identity or free-text
  column is a type error, never a coerced ``true``. Numeric cells are accepted only in the columns whose schema
  type is a number. Dates are valid only in the ``date`` column. A formula or a spreadsheet error cell is
  rejected rather than silently read as its cached value, and no formula is evaluated.
- **Value-free failures, always.** Every rejection carries source kind ``xlsx`` and an allowlisted column,
  ``record_type``, or integer count. No sheet name, path, cell text, unknown key, or library message appears,
  and a broad library exception family is normalized at one documented boundary with the raw cause suppressed.
"""

from __future__ import annotations

import os
from collections.abc import Iterator, Mapping
from datetime import date, datetime
from pathlib import Path
from typing import Any

from guiskindose.equipment_profiles.canonical import source_bytes_hash
from guiskindose.equipment_profiles.constants import (
    MAX_SOURCE_BYTES,
    SOURCE_KIND_XLSX,
    XLSX_SOURCE_SUFFIXES,
)
from guiskindose.equipment_profiles.csv_source import logical_rows, rows_to_document
from guiskindose.equipment_profiles.errors import EquipmentProfileError, ProfileErrorCode
from guiskindose.equipment_profiles.json_source import ParsedDocument
from guiskindose.equipment_profiles.tabular_common import ParsedRow, parse_row, validate_header
from guiskindose.equipment_profiles.validation import validate_document

#: openpyxl ``data_type`` tokens this adapter refuses rather than reading a cached value.
_FORMULA_DATA_TYPE = "f"
_ERROR_DATA_TYPE = "e"

#: The only flat column permitted to hold a spreadsheet date or datetime object.
_DATE_COLUMN = "date"

#: The only flat column permitted to hold a spreadsheet boolean.
_BOOLEAN_COLUMN = "promote_coordinate_wildcard"

#: Columns whose schema type is a number.
_NUMERIC_COLUMNS = frozenset({"kvp_kv", "filtration_added_mmcu", "filtration_added_mmal", "k_patient_support"})

#: The flat column whose type depends on the row's ``field`` value and is resolved in a second pass.
_DEFERRED_COLUMN = "value"

#: Columns whose schema type is text; a typed number or boolean there is a type error, never a coercion.
_TEXT_COLUMNS = frozenset(
    {
        "record_type",
        "profile_id",
        "manufacturer",
        "models",
        "manufacturer_aliases",
        "model_aliases",
        "source",
        "notes",
        "field",
        "device_model",
        "acquisition_plane",
        "evidence_status",
        "comment",
    }
)

_CELL_TEXT = "text"
_CELL_NUMBER = "number"
_CELL_BOOLEAN = "boolean"
_CELL_DATE = "date"
_CELL_DEFERRED = "deferred"

_COLUMN_KINDS: dict[str, str] = {
    **dict.fromkeys(_TEXT_COLUMNS, _CELL_TEXT),
    **dict.fromkeys(_NUMERIC_COLUMNS, _CELL_NUMBER),
    _BOOLEAN_COLUMN: _CELL_BOOLEAN,
    _DATE_COLUMN: _CELL_DATE,
    _DEFERRED_COLUMN: _CELL_DEFERRED,
}


def _read_bounded(raw: bytes, source_kind: str) -> None:
    if len(raw) > MAX_SOURCE_BYTES:
        raise EquipmentProfileError(
            ProfileErrorCode.OVERSIZED_SOURCE, source_kind=source_kind, count=MAX_SOURCE_BYTES
        )


def _bounded_read(path: Path, source_kind: str) -> bytes:
    """Read at most ``MAX_SOURCE_BYTES + 1`` bytes in one call, then apply the ceiling."""
    try:
        with path.open("rb") as handle:
            return handle.read(MAX_SOURCE_BYTES + 1)
    except (OSError, ValueError):
        raise EquipmentProfileError(ProfileErrorCode.UNREADABLE_SOURCE, source_kind=source_kind) from None


# ---------------------------------------------------------------------------
# Typed cell policy
# ---------------------------------------------------------------------------


def _cell_text(cell: Any, *, column: str | None, row_index: int, source_kind: str) -> str:
    """Convert one typed worksheet cell to the canonical flat text form, enforcing the column's type.

    ``column`` is the header name for the cell's position, or ``None`` while converting the header row (where no
    column name is known yet), in which case every cell must be text.
    """
    field_path = (column,) if column else ()
    data_type = getattr(cell, "data_type", None)
    if data_type == _FORMULA_DATA_TYPE:
        raise EquipmentProfileError(
            ProfileErrorCode.UNSAFE_FORMULA, source_kind=source_kind, field_path=field_path, row_index=row_index
        )
    if data_type == _ERROR_DATA_TYPE:
        raise EquipmentProfileError(
            ProfileErrorCode.SPREADSHEET_ERROR, source_kind=source_kind, field_path=field_path, row_index=row_index
        )
    value = getattr(cell, "value", cell)
    if value is None:
        return ""
    kind = _COLUMN_KINDS.get(column or "", _CELL_TEXT)
    return _typed_cell_text(value, kind=kind, field_path=field_path, row_index=row_index, source_kind=source_kind)


def _typed_cell_text(
    value: Any,
    *,
    kind: str,
    field_path: tuple[str, ...],
    row_index: int,
    source_kind: str,
) -> str:
    """Render a non-None, non-formula, non-error cell value under the column's declared type."""
    if kind == _CELL_DEFERRED:
        # The deferred column's type comes from the row's ``field`` and is resolved in a second pass, so this
        # pass stays permissive and :func:`_resolve_deferred` applies the authoritative rule.
        return _deferred_cell_text(value, field_path=field_path, row_index=row_index, source_kind=source_kind)
    if isinstance(value, bool):
        # Only the boolean column accepts a boolean cell. A boolean anywhere else is rejected with the code its
        # column would give a bad number or a bad text, never coerced into ``true`` or ``1``.
        if kind == _CELL_BOOLEAN:
            return "true" if value else "false"
        code = ProfileErrorCode.INVALID_NUMBER if kind == _CELL_NUMBER else ProfileErrorCode.INVALID_TYPE
        raise EquipmentProfileError(code, source_kind=source_kind, field_path=field_path, row_index=row_index)
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if kind == _CELL_BOOLEAN:
        raise EquipmentProfileError(
            ProfileErrorCode.INVALID_TYPE, source_kind=source_kind, field_path=field_path, row_index=row_index
        )
    if kind == _CELL_TEXT:
        return _text_cell_text(value, field_path=field_path, row_index=row_index, source_kind=source_kind)
    if kind == _CELL_NUMBER:
        return _number_cell_text(value, field_path=field_path, row_index=row_index, source_kind=source_kind)
    return _date_cell_text(value, field_path=field_path, row_index=row_index, source_kind=source_kind)


def _text_cell_text(
    value: Any,
    *,
    field_path: tuple[str, ...],
    row_index: int,
    source_kind: str,
) -> str:
    """Render a text-typed cell: numbers and booleans are rejected, not coerced."""
    if isinstance(value, str):
        return value
    raise EquipmentProfileError(
        ProfileErrorCode.INVALID_TYPE, source_kind=source_kind, field_path=field_path, row_index=row_index
    )


def _number_cell_text(
    value: Any,
    *,
    field_path: tuple[str, ...],
    row_index: int,
    source_kind: str,
) -> str:
    """Render a number-typed cell. Booleans are numbers in Excel; here they never are."""
    if isinstance(value, bool):
        raise EquipmentProfileError(
            ProfileErrorCode.INVALID_NUMBER, source_kind=source_kind, field_path=field_path, row_index=row_index
        )
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return _finite_float_text(value, field_path=field_path, row_index=row_index, source_kind=source_kind)
    raise EquipmentProfileError(
        ProfileErrorCode.INVALID_NUMBER, source_kind=source_kind, field_path=field_path, row_index=row_index
    )


def _finite_float_text(
    value: float,
    *,
    field_path: tuple[str, ...],
    row_index: int,
    source_kind: str,
) -> str:
    if value != value or value in (float("inf"), float("-inf")):
        raise EquipmentProfileError(
            ProfileErrorCode.INVALID_NUMBER, source_kind=source_kind, field_path=field_path, row_index=row_index
        )
    return repr(value)


def _date_cell_text(
    value: Any,
    *,
    field_path: tuple[str, ...],
    row_index: int,
    source_kind: str,
) -> str:
    """Render a ``date`` cell: a typed date must be exactly midnight, and a string is passed through.

    A string reaches the shared flat validator, which enforces strict ``YYYY-MM-DD``, so a text date authored
    the same way as the CSV form stays equivalent. Only a *typed* date or datetime object is checked here, and
    it must name exactly midnight of one day.
    """
    if isinstance(value, str):
        return value
    if not isinstance(value, (datetime, date)):
        raise EquipmentProfileError(
            ProfileErrorCode.INVALID_DATE_OBJECT, source_kind=source_kind, field_path=field_path, row_index=row_index
        )
    if isinstance(value, datetime) and (
        value.tzinfo is not None
        or (value.hour, value.minute, value.second, value.microsecond) != (0, 0, 0, 0)
    ):
        raise EquipmentProfileError(
            ProfileErrorCode.INVALID_DATE_OBJECT, source_kind=source_kind, field_path=field_path, row_index=row_index
        )
    return f"{value.year:04d}-{value.month:02d}-{value.day:02d}"


def _deferred_cell_text(
    value: Any,
    *,
    field_path: tuple[str, ...],
    row_index: int,
    source_kind: str,
) -> str:
    """Render the ``value`` column permissively; its type is resolved in a second pass.

    Text, numbers, and booleans are all carried here and re-checked against the resolved coordinate leaf in
    :func:`_resolve_deferred`. Dates are never valid, because no coordinate leaf is a date.
    """
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date)):
        raise EquipmentProfileError(
            ProfileErrorCode.INVALID_DATE_OBJECT, source_kind=source_kind, field_path=field_path, row_index=row_index
        )
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return _finite_float_text(value, field_path=field_path, row_index=row_index, source_kind=source_kind)
    if isinstance(value, str):
        return value
    raise EquipmentProfileError(
        ProfileErrorCode.INVALID_TYPE, source_kind=source_kind, field_path=field_path, row_index=row_index
    )


# ---------------------------------------------------------------------------
# Row conversion
# ---------------------------------------------------------------------------


def _row_cells(
    row: Any,
    names: tuple[str, ...] | None,
    source_kind: str,
    row_index: int,
) -> list[str]:
    """Convert one worksheet row to the canonical flat text cells, never truncating a populated cell."""
    values = list(row)
    if names is None:
        return _header_cells(values, source_kind, row_index)
    return _body_cells(values, names, source_kind, row_index)


def _header_cells(values: list[Any], source_kind: str, row_index: int) -> list[str]:
    """Convert a header row, dropping openpyxl's trailing padding but no populated cell.

    A populated cell beyond the header is kept, so ``validate_header`` rejects it as an unknown column name
    rather than it being silently dropped. A formula or error cell in any position raises before that.
    """
    cells = [_cell_text(value, column=None, row_index=row_index, source_kind=source_kind) for value in values]
    while cells and not cells[-1]:
        cells.pop()
    return cells


def _body_cells(values: list[Any], names: tuple[str, ...], source_kind: str, row_index: int) -> list[str]:
    """Convert a body row, rejecting populated cells beyond the header instead of dropping them."""
    cells = [
        _cell_text(value, column=name, row_index=row_index, source_kind=source_kind)
        for name, value in zip(names, values, strict=False)
    ]
    if len(values) > len(names):
        for index in range(len(names), len(values)):
            if not _is_blank(values[index]):
                raise EquipmentProfileError(
                    ProfileErrorCode.UNKNOWN_FIELD, source_kind=source_kind, row_index=row_index, count=index
                )
    while len(cells) < len(names):
        cells.append("")
    return _resolve_deferred(cells, values, names, source_kind, row_index)


def _cell_value(cell: Any) -> Any:
    """Return the underlying value of a worksheet cell, or the value itself when it is already raw."""
    if cell is None:
        return None
    value = getattr(cell, "value", None)
    return cell if value is None and not hasattr(cell, "value") else value


def _resolve_deferred(
    cells: list[str],
    values: list[Any],
    names: tuple[str, ...],
    source_kind: str,
    row_index: int,
) -> list[str]:
    """Re-convert the ``value`` cell under the coordinate leaf this row names.

    The ``value`` column's accepted type depends on the row's ``field``, so a number, boolean, or date is
    accepted or rejected by the leaf it addresses, not by a generic rule.
    """
    if "record_type" not in names or "field" not in names or "value" not in names:
        return cells
    if cells[names.index("record_type")] != "coordinate":
        return cells
    field = cells[names.index("field")]
    if not field:
        return cells
    kind = _coordinate_leaf_kind(field, source_kind, row_index)
    position = names.index("value")
    raw = _cell_value(values[position]) if position < len(values) else None
    cells[position] = _typed_cell_text(
        raw, kind=kind, field_path=("value",), row_index=row_index, source_kind=source_kind
    )
    return cells


def _coordinate_leaf_kind(field: str, source_kind: str, row_index: int) -> str:
    """Return the cell kind the leaf named by ``field`` accepts."""
    from guiskindose.equipment_profiles.tabular_common import (
        SWAP_LATERAL_LONGITUDINAL,
        TRANSLATION_OFFSET,
        coordinate_leaf_path,
    )

    block, _leaf = coordinate_leaf_path(source_kind, row_index, field)
    if block in {TRANSLATION_OFFSET, "detector_side_length"}:
        return _CELL_NUMBER
    if block == SWAP_LATERAL_LONGITUDINAL:
        return _CELL_BOOLEAN
    return _CELL_TEXT


def _is_blank(cell: Any) -> bool:
    """True for a cell openpyxl reports as empty: absent, or a cell whose value is blank."""
    value = _cell_value(cell)
    if value is None:
        return True
    return isinstance(value, str) and value == ""


def _iter_rows(worksheet: Any) -> Iterator[Any]:
    """Bound reader-emitted rows, including padding, after XML coordinate validation."""
    from guiskindose.equipment_profiles import constants

    cells = 0
    for index, row in enumerate(worksheet.iter_rows(), start=1):
        cells += len(row)
        if (index > constants.MAX_XLSX_PREFLIGHT_ROWS
                or len(row) > constants.MAX_XLSX_PREFLIGHT_COLS
                or cells > constants.MAX_XLSX_PREFLIGHT_CELLS):
            raise EquipmentProfileError(ProfileErrorCode.ROW_LIMIT_EXCEEDED, source_kind=SOURCE_KIND_XLSX)
        yield row


# ---------------------------------------------------------------------------
# Workbook access
# ---------------------------------------------------------------------------


def _open_workbook(raw: bytes, source_kind: str) -> Any:
    """Open the workbook from the byte snapshot with safe, non-materializing options."""
    import io

    from openpyxl import load_workbook

    try:
        return load_workbook(
            io.BytesIO(raw),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception:
        # openpyxl raises a wide family of exceptions for a malformed container. They all collapse to one fixed
        # code. This broad normalization is confined to the adapter boundary, and the raw library message is
        # never included in the error.
        raise EquipmentProfileError(
            ProfileErrorCode.MALFORMED_CONTAINER, source_kind=source_kind
        ) from None


def _close_workbook(workbook: Any) -> None:
    """Close the workbook reliably, swallowing only a close-time failure.

    A failure while closing must never mask a successful parse, so the close is isolated. Only the close call
    is suppressed, and no exception object is inspected or logged.
    """
    import contextlib

    with contextlib.suppress(Exception):
        workbook.close()


def _worksheets(workbook: Any) -> Any:
    return workbook.worksheets


# ---------------------------------------------------------------------------
# Document assembly
# ---------------------------------------------------------------------------


def _check_catalog_parts(raw: bytes, worksheets: Any, source_kind: str) -> None:
    """Require every reader-selected worksheet to have passed XML preflight."""
    import io
    import zipfile

    from guiskindose.equipment_profiles.xlsx_xml import worksheet_entry_names

    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        validated = set(worksheet_entry_names(archive))
    for worksheet in worksheets:
        if getattr(worksheet, "_worksheet_path", None) not in validated:
            raise EquipmentProfileError(ProfileErrorCode.MALFORMED_CONTAINER, source_kind=source_kind)


def _rows(raw: bytes, *, selected: int | None, source_kind: str) -> Iterator[ParsedRow]:
    """Preflight the container and its XML, then stream validated flat rows.

    The container inflation, the container's structure, and every worksheet part's actual coordinates are
    validated **before** the reader opens, so a forged declared dimension cannot cause either a silent
    truncation or an oversized allocation.
    """
    from guiskindose.equipment_profiles.xlsx_preflight import (
        preflight_container,
        resolve_sheet,
    )
    from guiskindose.equipment_profiles.xlsx_xml import validate_worksheet_parts

    preflight_container(raw, source_kind=source_kind)
    validate_worksheet_parts(raw, source_kind=source_kind)

    workbook = _open_workbook(raw, source_kind)
    try:
        worksheets = _worksheets(workbook)
        count = len(worksheets)
        index = resolve_sheet(count, selected, source_kind=source_kind)
        _check_catalog_parts(raw, worksheets, source_kind)
        sheet = worksheets[index]
        # The declared dimension was deliberately not trusted for sizing: discard it so the reader yields the
        # real rows validated above rather than only the ones it claims exist.
        reset = getattr(sheet, "reset_dimensions", None)
        if callable(reset):
            reset()
        header: dict[str, int] | None = None
        data_rows = 0
        for row in _iter_rows(sheet):
            if header is None:
                names = tuple(_row_cells(row, None, source_kind, data_rows))
                header = validate_header(list(names), source_kind)
                continue
            yield parse_row(
                _row_cells(row, _ordered_names(header), source_kind, data_rows),
                header,
                source_kind,
                source_index=data_rows,
            )
            data_rows += 1
    finally:
        _close_workbook(workbook)


def _ordered_names(positions: Mapping[str, int]) -> tuple[str, ...]:
    """Return the header names ordered by their position."""
    ordered = sorted(positions, key=positions.__getitem__)
    return tuple(ordered)


def _guarded_rows(raw: bytes, *, selected: int | None, source_kind: str) -> Iterator[ParsedRow]:
    """Boundary over the whole untrusted-container path: preflight, reader, and iteration.

    Every library family this path can raise is normalized to one fixed code with the raw cause suppressed.
    The boundary encloses *only* untrusted container parsing, and ``EquipmentProfileError`` passes through
    untouched so a specific code is never masked.
    """
    try:
        yield from _rows(raw, selected=selected, source_kind=source_kind)
    except EquipmentProfileError:
        raise
    except Exception:
        raise EquipmentProfileError(ProfileErrorCode.MALFORMED_CONTAINER, source_kind=source_kind) from None


def parse_xlsx_bytes(raw: bytes, *, sheet: int | None = None) -> ParsedDocument:
    """Parse and validate an XLSX profile document from raw bytes.

    ``sheet`` is the zero-based sheet index to read. A single-sheet workbook defaults to sheet ``0``; a
    workbook with two or more sheets requires an explicit selection, otherwise ``sheet_required``.
    """
    if not raw:
        raise EquipmentProfileError(ProfileErrorCode.EMPTY_SOURCE, source_kind=SOURCE_KIND_XLSX)
    _read_bounded(raw, SOURCE_KIND_XLSX)
    document = rows_to_document(
        _guarded_rows(raw, selected=sheet, source_kind=SOURCE_KIND_XLSX), SOURCE_KIND_XLSX
    )
    return ParsedDocument(
        profiles=validate_document(document, source_kind=SOURCE_KIND_XLSX),
        source_kind=SOURCE_KIND_XLSX,
        source_hash=source_bytes_hash(raw),
        logical_rows=logical_rows(document),
    )


def parse_xlsx_file(path: str | os.PathLike[str], *, sheet: int | None = None) -> ParsedDocument:
    """Read, parse, and validate an XLSX profile file.

    The path never appears in an error. The file is read with one bounded read of at most
    ``MAX_SOURCE_BYTES + 1`` bytes, so an oversized source is rejected from the read length itself.
    """
    candidate = Path(path)
    if candidate.suffix.lower() not in XLSX_SOURCE_SUFFIXES:
        raise EquipmentProfileError(
            ProfileErrorCode.UNSUPPORTED_SOURCE_SUFFIX, source_kind=SOURCE_KIND_XLSX
        )
    raw = _bounded_read(candidate, SOURCE_KIND_XLSX)
    _read_bounded(raw, SOURCE_KIND_XLSX)
    return parse_xlsx_bytes(raw, sheet=sheet)
